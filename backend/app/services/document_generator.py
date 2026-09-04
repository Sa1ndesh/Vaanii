from docx import Document
import io

from app.services.ollama_client import chat

# System instruction for the AI document drafter
SYSTEM_INSTRUCTION = """You are an expert AI legal assistant specializing in drafting clear, professional, and accurate legal documents based on Indian law.
-   **Task:** Generate the full text for the requested legal document.
-   **Input:** You will receive the document type and key details (like party names, dates, terms).
-   **Output:** Return ONLY the complete, professionally formatted text of the document.
-   **Formatting:** Use standard legal document formatting. Use newlines (\\n) for paragraphs and numbering.
-   **Tone:** Formal, precise, and authoritative.
-   **CRITICAL:** Do NOT include any conversational text, disclaimers, or explanations. Only output the raw document text itself.
"""

# --- Main Service Function ---

async def create_document(doc_type: str, details: dict) -> io.BytesIO:
    """
    Generates a legal document text using Ollama and creates a .docx file in memory.
    """
    # 1. Create a detailed prompt for the AI
    prompt = f"Generate a {doc_type} with the following details:\n"
    for key, value in details.items():
        if value:  # Only include details that are filled in
            prompt += f"- {key.replace('_', ' ').title()}: {value}\n"

    prompt += "\nReturn only the full, formatted text of the legal document, ready for a .docx file."

    try:
        # 2. Get the raw document text from Ollama
        print(f"Generating draft for: {doc_type}")
        document_text = await chat(
            prompt,
            system=SYSTEM_INSTRUCTION,
            temperature=0.5,
            max_tokens=4096,
        )
        print("Draft received from Ollama.")

        # 3. Create a .docx file in memory
        document = Document()

        # Add a title
        document.add_heading(doc_type, level=0)

        # Add the AI-generated text (preserve paragraph breaks)
        for para in document_text.split('\n'):
            if para.strip():
                document.add_paragraph(para)

        # 4. Save the document to a byte stream
        file_stream = io.BytesIO()
        document.save(file_stream)
        file_stream.seek(0)

        print(f"Successfully created .docx for {doc_type} in memory.")
        return file_stream

    except Exception as e:
        print(f"Error in create_document service: {e}")
        raise